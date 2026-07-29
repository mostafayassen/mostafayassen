"""Excel (.xlsx) import/export for tasks and reports, using openpyxl.

This is the one integration with no auth/network dependency, so it is fully
self-contained and should work reliably offline.

Import format (first row = header, columns matched by name, case-insensitive,
order doesn't matter -- any missing optional columns just fall back to
defaults):

    Title | Description | Project | Tags | Quadrant | Deadline |
    Estimate Minutes | Status

Only "Title" is required. "Tags" is a comma-separated list.  "Quadrant" may
be one of: urgent_important, not_urgent_important, urgent_not_important,
not_urgent_not_important (or the human labels from models.QUADRANT_LABELS).
"""

from __future__ import annotations

from datetime import datetime
from typing import List, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from app.models import ALL_QUADRANTS, QUADRANT_LABELS, Task

TASK_EXPORT_HEADERS = [
    "Title",
    "Description",
    "Project",
    "Tags",
    "Quadrant",
    "Deadline",
    "Estimate Minutes",
    "Actual Minutes",
    "Status",
    "Source",
]

_LABEL_TO_QUADRANT = {v.lower(): k for k, v in QUADRANT_LABELS.items()}


def _normalize_quadrant(value: Optional[str]) -> str:
    if not value:
        return ALL_QUADRANTS[1]  # default: not_urgent_important
    value = str(value).strip()
    lowered = value.lower()
    if lowered in ALL_QUADRANTS:
        return lowered
    if lowered in _LABEL_TO_QUADRANT:
        return _LABEL_TO_QUADRANT[lowered]
    return ALL_QUADRANTS[1]


def _normalize_deadline(value) -> Optional[str]:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.isoformat(timespec="seconds")
    # Try a couple of common string formats; fall back to storing raw text.
    text = str(value).strip()
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, fmt).isoformat(timespec="seconds")
        except ValueError:
            continue
    try:
        return datetime.fromisoformat(text).isoformat(timespec="seconds")
    except ValueError:
        return text  # give up gracefully, store as-is


def import_tasks_from_xlsx(path: str, sheet_name: Optional[str] = None) -> List[Task]:
    """Read a task list from an existing .xlsx file.

    Raises FileNotFoundError / openpyxl exceptions on bad input -- callers
    (UI code) are expected to catch and show a friendly message.
    """
    wb = load_workbook(path, data_only=True)
    ws: Worksheet = wb[sheet_name] if sheet_name else wb.active

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    header = [str(h).strip().lower() if h is not None else "" for h in rows[0]]

    def col(*names: str) -> Optional[int]:
        for name in names:
            if name in header:
                return header.index(name)
        return None

    idx_title = col("title", "task", "name")
    idx_desc = col("description", "notes", "details")
    idx_project = col("project", "category")
    idx_tags = col("tags", "tag")
    idx_quadrant = col("quadrant", "priority")
    idx_deadline = col("deadline", "due date", "due")
    idx_estimate = col("estimate minutes", "estimate", "estimated minutes")
    idx_status = col("status")
    idx_source = col("source")

    if idx_title is None:
        raise ValueError(
            "The Excel sheet must have a 'Title' column (row 1) to import tasks."
        )

    tasks: List[Task] = []
    for raw_row in rows[1:]:
        if raw_row is None or all(cell in (None, "") for cell in raw_row):
            continue
        title = raw_row[idx_title] if idx_title < len(raw_row) else None
        if not title:
            continue

        def get(idx: Optional[int]):
            if idx is None or idx >= len(raw_row):
                return None
            return raw_row[idx]

        tags_raw = get(idx_tags)
        tags = (
            [t.strip() for t in str(tags_raw).split(",") if t.strip()]
            if tags_raw
            else []
        )

        estimate_raw = get(idx_estimate)
        try:
            estimate_minutes = int(estimate_raw) if estimate_raw not in (None, "") else 0
        except (ValueError, TypeError):
            estimate_minutes = 0

        task = Task(
            title=str(title).strip(),
            description=str(get(idx_desc) or ""),
            project=str(get(idx_project) or ""),
            tags=tags,
            quadrant=_normalize_quadrant(get(idx_quadrant)),
            deadline=_normalize_deadline(get(idx_deadline)),
            estimate_minutes=estimate_minutes,
            status=str(get(idx_status) or "todo").strip().lower() or "todo",
            source=str(get(idx_source) or "excel"),
        )
        tasks.append(task)

    return tasks


def export_tasks_to_xlsx(tasks: List[Task], path: str, sheet_title: str = "Tasks") -> str:
    """Write the given tasks out to a new .xlsx file. Returns the path written."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title

    ws.append(TASK_EXPORT_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for t in tasks:
        ws.append(
            [
                t.title,
                t.description,
                t.project,
                ", ".join(t.tags or []),
                QUADRANT_LABELS.get(t.quadrant, t.quadrant),
                t.deadline or "",
                t.estimate_minutes,
                t.actual_minutes,
                t.status,
                t.source,
            ]
        )

    _autosize_columns(ws)
    wb.save(path)
    return path


def export_report_to_xlsx(report: dict, path: str) -> str:
    """Write a weekly report dict (see reports_view.build_weekly_report) to xlsx.

    Produces a workbook with three sheets: Summary, Time By Project,
    Overdue Tasks.
    """
    wb = Workbook()

    summary_ws = wb.active
    summary_ws.title = "Summary"
    summary_ws.append(["Metric", "Value"])
    for cell in summary_ws[1]:
        cell.font = Font(bold=True)
    for key, value in report.get("summary", {}).items():
        summary_ws.append([key, value])
    _autosize_columns(summary_ws)

    time_ws = wb.create_sheet("Time By Project")
    time_ws.append(["Project", "Minutes", "Hours"])
    for cell in time_ws[1]:
        cell.font = Font(bold=True)
    for project, minutes in report.get("time_by_project", {}).items():
        time_ws.append([project or "(no project)", minutes, round(minutes / 60, 2)])
    _autosize_columns(time_ws)

    overdue_ws = wb.create_sheet("Overdue Tasks")
    overdue_ws.append(["Title", "Project", "Deadline", "Status"])
    for cell in overdue_ws[1]:
        cell.font = Font(bold=True)
    for t in report.get("overdue_tasks", []):
        overdue_ws.append([t.title, t.project, t.deadline, t.status])
    _autosize_columns(overdue_ws)

    wb.save(path)
    return path


def _autosize_columns(ws: Worksheet, min_width: int = 10, max_width: int = 50) -> None:
    for column_cells in ws.columns:
        length = max((len(str(c.value)) for c in column_cells if c.value is not None), default=0)
        col_letter = column_cells[0].column_letter
        ws.column_dimensions[col_letter].width = max(min_width, min(max_width, length + 2))
