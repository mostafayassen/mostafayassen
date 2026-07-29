"""Unit tests for app/integrations/excel_io.py -- a full round-trip of
export -> import, plus edge cases around header matching and defaults.
No network or Windows-only dependency, so these run anywhere.
"""

from __future__ import annotations

import os

from app.integrations import excel_io
from app.models import QUADRANT_URGENT_IMPORTANT, Task


def _sample_tasks():
    return [
        Task(
            title="Prepare board deck",
            description="Slides for Q3 review",
            project="Leadership",
            tags=["deck", "quarterly"],
            quadrant=QUADRANT_URGENT_IMPORTANT,
            deadline="2026-08-01T09:00:00",
            estimate_minutes=120,
            actual_minutes=30,
            status="in_progress",
            source="manual",
        ),
        Task(
            title="Reply to vendor email",
            description="",
            project="Ops",
            tags=[],
            quadrant="urgent_not_important",
            deadline=None,
            estimate_minutes=15,
            actual_minutes=0,
            status="todo",
            source="email",
        ),
    ]


def test_export_then_import_round_trip(tmp_path):
    path = str(tmp_path / "tasks.xlsx")
    tasks = _sample_tasks()

    written_path = excel_io.export_tasks_to_xlsx(tasks, path)
    assert os.path.exists(written_path)

    imported = excel_io.import_tasks_from_xlsx(path)
    assert len(imported) == 2

    titles = {t.title for t in imported}
    assert titles == {"Prepare board deck", "Reply to vendor email"}

    deck = next(t for t in imported if t.title == "Prepare board deck")
    assert deck.project == "Leadership"
    assert set(deck.tags) == {"deck", "quarterly"}
    assert deck.estimate_minutes == 120
    assert deck.quadrant == QUADRANT_URGENT_IMPORTANT


def test_import_requires_title_column(tmp_path):
    import openpyxl

    path = str(tmp_path / "bad.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Description", "Project"])
    ws.append(["no title here", "Ops"])
    wb.save(path)

    try:
        excel_io.import_tasks_from_xlsx(path)
        assert False, "expected ValueError for missing Title column"
    except ValueError as exc:
        assert "Title" in str(exc)


def test_import_skips_blank_rows_and_handles_missing_optional_columns(tmp_path):
    import openpyxl

    path = str(tmp_path / "minimal.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Title"])
    ws.append(["Just a title"])
    ws.append([None])  # blank row should be skipped
    wb.save(path)

    tasks = excel_io.import_tasks_from_xlsx(path)
    assert len(tasks) == 1
    assert tasks[0].title == "Just a title"
    assert tasks[0].estimate_minutes == 0
    assert tasks[0].tags == []


def test_import_handles_quadrant_human_labels(tmp_path):
    import openpyxl

    path = str(tmp_path / "labels.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Title", "Quadrant"])
    ws.append(["Urgent thing", "Urgent & Important (Do First)"])
    wb.save(path)

    tasks = excel_io.import_tasks_from_xlsx(path)
    assert tasks[0].quadrant == QUADRANT_URGENT_IMPORTANT


def test_export_report_to_xlsx(tmp_path):
    # Deliberately does not import app.ui.reports_view here: excel_io must
    # stay fully independent of the PySide6 GUI stack (see module docstring
    # and README) so these tests run in headless/CI environments too.
    report = {
        "summary": {"Total time logged (minutes)": 90, "Tasks completed": 2},
        "time_by_project": {"Ops": 45, "Leadership": 45},
        "overdue_tasks": [],
    }
    path = str(tmp_path / "report.xlsx")
    written = excel_io.export_report_to_xlsx(report, path)
    assert os.path.exists(written)

    import openpyxl

    wb = openpyxl.load_workbook(written)
    assert set(wb.sheetnames) == {"Summary", "Time By Project", "Overdue Tasks"}
